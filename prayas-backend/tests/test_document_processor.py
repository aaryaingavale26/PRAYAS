import io
import sys
import unittest
from pathlib import Path

# Add backend directory to sys.path
BACKEND_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND_DIR))

import docx
import pymupdf
from app.services.document_processor import (
    DocumentProcessingError,
    extract_text_from_docx,
    extract_text_from_pdf,
)


class TestDocumentProcessor(unittest.TestCase):
    def _create_sample_pdf(self, text: str = "") -> bytes:
        """Helper to create an in-memory PDF."""
        doc = pymupdf.open()
        page = doc.new_page(width=595, height=842)
        if text:
            page.insert_text((72, 72), text)
        pdf_bytes = doc.tobytes()
        doc.close()
        return pdf_bytes

    def _create_sample_docx(self, paragraphs=None, tables=None) -> bytes:
        """Helper to create an in-memory DOCX."""
        doc = docx.Document()
        if paragraphs:
            for p in paragraphs:
                doc.add_paragraph(p)
        if tables:
            for table_data in tables:
                table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                for r_idx, row in enumerate(table_data):
                    for c_idx, val in enumerate(row):
                        table.cell(r_idx, c_idx).text = val
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    # 1. Valid PDF containing text
    def test_extract_text_from_valid_pdf(self):
        sample_text = "John Doe\nSoftware Engineer Resume\nSkills: Python, FastAPI"
        pdf_bytes = self._create_sample_pdf(sample_text)
        extracted = extract_text_from_pdf(pdf_bytes)
        self.assertIn("John Doe", extracted)
        self.assertIn("FastAPI", extracted)

    # 2. Valid DOCX containing text
    def test_extract_text_from_valid_docx(self):
        sample_paragraphs = ["Jane Smith", "Job Application - Accessibility Advocate"]
        docx_bytes = self._create_sample_docx(paragraphs=sample_paragraphs)
        extracted = extract_text_from_docx(docx_bytes)
        self.assertIn("Jane Smith", extracted)
        self.assertIn("Accessibility Advocate", extracted)

    # 3. DOCX containing a table
    def test_extract_text_from_docx_with_table(self):
        table_data = [
            ["Skill", "Proficiency"],
            ["Python", "Expert"],
            ["Screen Readers", "Advanced"],
        ]
        docx_bytes = self._create_sample_docx(paragraphs=["Skills Summary"], tables=[table_data])
        extracted = extract_text_from_docx(docx_bytes)
        self.assertIn("Skills Summary", extracted)
        self.assertIn("Python", extracted)
        self.assertIn("Screen Readers", extracted)
        self.assertIn("Advanced", extracted)

    # 4. Valid blank PDF
    def test_extract_text_from_blank_pdf(self):
        blank_pdf_bytes = self._create_sample_pdf(text="")
        extracted = extract_text_from_pdf(blank_pdf_bytes)
        self.assertEqual(extracted, "")

    # 5. Valid blank DOCX
    def test_extract_text_from_blank_docx(self):
        blank_docx_bytes = self._create_sample_docx()
        extracted = extract_text_from_docx(blank_docx_bytes)
        self.assertEqual(extracted, "")

    # 6. Corrupted PDF
    def test_corrupted_pdf_raises_error(self):
        corrupted_bytes = b"%PDF-1.4 completely invalid and corrupted random binary content \x00\xff\xfe"
        with self.assertRaises(DocumentProcessingError) as ctx:
            extract_text_from_pdf(corrupted_bytes)
        self.assertIn("Invalid or corrupted file", str(ctx.exception))

    # 7. Corrupted DOCX
    def test_corrupted_docx_raises_error(self):
        corrupted_bytes = b"PK\x03\x04 fake corrupted zip header content not a real docx"
        with self.assertRaises(DocumentProcessingError) as ctx:
            extract_text_from_docx(corrupted_bytes)
        self.assertIn("Invalid or corrupted file", str(ctx.exception))

    # 8. Empty input bytes
    def test_empty_bytes_raises_error(self):
        with self.assertRaises(DocumentProcessingError) as ctx_pdf:
            extract_text_from_pdf(b"")
        self.assertIn("empty", str(ctx_pdf.exception).lower())

        with self.assertRaises(DocumentProcessingError) as ctx_docx:
            extract_text_from_docx(b"")
        self.assertIn("empty", str(ctx_docx.exception).lower())


if __name__ == "__main__":
    unittest.main()
