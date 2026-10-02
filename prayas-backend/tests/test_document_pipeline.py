import io
import os
import sys
import unittest
import uuid
from pathlib import Path
from unittest.mock import MagicMock, patch

# Ensure backend directory is in sys.path
BACKEND_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND_DIR))

import docx
import pymupdf
from starlette.testclient import TestClient

from app.core.auth import AuthenticatedUser, get_current_user
from app.main import app
from app.services.document_metadata import DocumentMetadataError
from app.services.document_processor import DocumentProcessingError
from app.services.file_validator import DOCX_MIME, PDF_MIME
from app.services.storage_service import StorageServiceError


class TestDocumentPipelineIntegration(unittest.TestCase):
    """
    Integration tests covering the end-to-end document upload and retrieval workflows.
    Verifies that file validation, text extraction, storage management, and database metadata
    interact correctly without leaks or orphaned resources.
    """

    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)

        # Generate a valid PDF document with realistic content
        pdf_doc = pymupdf.open()
        page = pdf_doc.new_page(width=595, height=842)
        page.insert_text((72, 72), "Alex Morgan\nAccessible Frontend Developer\nExperience: React, WCAG 2.1")
        cls.valid_pdf_bytes = pdf_doc.tobytes()
        pdf_doc.close()

        # Generate a valid DOCX document with text and table
        word_doc = docx.Document()
        word_doc.add_paragraph("Jordan Taylor\nAccessibility Specialist Resume")
        table = word_doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Certification"
        table.cell(0, 1).text = "CPACC"
        table.cell(1, 0).text = "Year"
        table.cell(1, 1).text = "2025"
        buf = io.BytesIO()
        word_doc.save(buf)
        cls.valid_docx_bytes = buf.getvalue()

        cls.user_id = "11111111-2222-3333-4444-555555555555"
        app.dependency_overrides[get_current_user] = lambda: AuthenticatedUser(user_id=cls.user_id)

    @classmethod
    def tearDownClass(cls):
        app.dependency_overrides.pop(get_current_user, None)

    # =========================================================================
    # Task 2: Upload Pipeline Integration Tests
    # =========================================================================

    # 1. Successful PDF Upload
    @patch("app.api.documents.upload_document")
    @patch("app.api.documents.save_document_metadata")
    def test_pipeline_successful_pdf_upload(self, mock_save_meta, mock_upload):
        mock_upload.return_value = {"storage_path": "path", "status": "uploaded"}
        mock_save_meta.side_effect = lambda **kwargs: {
            "id": kwargs.get("document_id", str(uuid.uuid4())),
            "user_id": kwargs["user_id"],
            "file_name": kwargs["file_name"],
            "storage_path": kwargs["storage_path"],
            "file_type": kwargs["file_type"],
            "file_size": kwargs["file_size"],
        }

        response = self.client.post(
            "/api/v1/documents/upload",
            data={"user_id": self.user_id},
            files={"file": ("alex_resume.pdf", self.valid_pdf_bytes, PDF_MIME)},
        )

        self.assertEqual(response.status_code, 200)
        data = response.json()

        # Verify response structure
        self.assertEqual(data["file_name"], "alex_resume.pdf")
        self.assertEqual(data["file_type"], PDF_MIME)
        self.assertEqual(data["file_size"], len(self.valid_pdf_bytes))
        self.assertEqual(data["status"], "uploaded")
        self.assertIn("Alex Morgan", data["extracted_text"])
        self.assertIn("WCAG 2.1", data["extracted_text"])

        # Verify Storage upload called with correct bytes and content type
        self.assertTrue(mock_upload.called)
        upload_call = mock_upload.call_args[1]
        self.assertEqual(upload_call["file_bytes"], self.valid_pdf_bytes)
        self.assertEqual(upload_call["content_type"], PDF_MIME)
        self.assertTrue(upload_call["storage_path"].startswith(f"{self.user_id}/"))

        # Verify Metadata saved
        self.assertTrue(mock_save_meta.called)
        save_call = mock_save_meta.call_args[1]
        self.assertEqual(save_call["user_id"], self.user_id)
        self.assertEqual(save_call["file_name"], "alex_resume.pdf")

    # 2. Successful DOCX Upload
    @patch("app.api.documents.upload_document")
    @patch("app.api.documents.save_document_metadata")
    def test_pipeline_successful_docx_upload(self, mock_save_meta, mock_upload):
        mock_upload.return_value = {"storage_path": "path", "status": "uploaded"}
        mock_save_meta.side_effect = lambda **kwargs: {
            "id": kwargs.get("document_id", str(uuid.uuid4())),
            "user_id": kwargs["user_id"],
            "file_name": kwargs["file_name"],
            "storage_path": kwargs["storage_path"],
            "file_type": kwargs["file_type"],
            "file_size": kwargs["file_size"],
        }

        response = self.client.post(
            "/api/v1/documents/upload",
            data={"user_id": self.user_id},
            files={"file": ("jordan_resume.docx", self.valid_docx_bytes, DOCX_MIME)},
        )

        self.assertEqual(response.status_code, 200)
        data = response.json()

        self.assertEqual(data["file_name"], "jordan_resume.docx")
        self.assertEqual(data["file_type"], DOCX_MIME)
        self.assertEqual(data["file_size"], len(self.valid_docx_bytes))
        self.assertIn("Jordan Taylor", data["extracted_text"])
        self.assertIn("CPACC", data["extracted_text"])

        self.assertTrue(mock_upload.called)
        self.assertTrue(mock_save_meta.called)

    # 3. Invalid File (Rejected early without storage or database calls)
    @patch("app.api.documents.upload_document")
    @patch("app.api.documents.save_document_metadata")
    def test_pipeline_invalid_file_rejected_early(self, mock_save_meta, mock_upload):
        # Invalid extension and magic bytes
        response = self.client.post(
            "/api/v1/documents/upload",
            data={"user_id": self.user_id},
            files={"file": ("unsupported.exe", b"MZbinaryexecutable", "application/octet-stream")},
        )

        self.assertEqual(response.status_code, 400)
        # Ensure downstream storage and DB are NEVER called
        self.assertFalse(mock_upload.called)
        self.assertFalse(mock_save_meta.called)

    # 4. Metadata Save Failure (Triggers Storage Rollback / Cleanup)
    @patch("app.api.documents.upload_document")
    @patch("app.api.documents.save_document_metadata")
    @patch("app.api.documents.delete_document")
    def test_pipeline_metadata_failure_triggers_storage_rollback(self, mock_delete, mock_save_meta, mock_upload):
        mock_upload.return_value = {"status": "uploaded"}
        mock_save_meta.side_effect = DocumentMetadataError("Database connection lost during insert")

        response = self.client.post(
            "/api/v1/documents/upload",
            data={"user_id": self.user_id},
            files={"file": ("resume.pdf", self.valid_pdf_bytes, PDF_MIME)},
        )

        self.assertEqual(response.status_code, 500)
        # Verify rollback was called to eliminate orphaned storage file
        self.assertTrue(mock_delete.called)
        cleaned_path = mock_delete.call_args[0][0]
        self.assertTrue(cleaned_path.startswith(f"{self.user_id}/"))

    # 5. Storage Upload Failure (Aborts before Database Metadata insert)
    @patch("app.api.documents.upload_document")
    @patch("app.api.documents.save_document_metadata")
    def test_pipeline_storage_failure_aborts_before_metadata_save(self, mock_save_meta, mock_upload):
        mock_upload.side_effect = StorageServiceError("Supabase Storage bucket upload quota exceeded")

        response = self.client.post(
            "/api/v1/documents/upload",
            data={"user_id": self.user_id},
            files={"file": ("resume.pdf", self.valid_pdf_bytes, PDF_MIME)},
        )

        self.assertEqual(response.status_code, 502)
        self.assertIn("Storage upload failed", response.json()["error"])
        # Ensure metadata save is NEVER called when upload fails
        self.assertFalse(mock_save_meta.called)

    # =========================================================================
    # Task 3: Retrieval Pipeline Integration Tests
    # =========================================================================

    # 1. Successful PDF Retrieval
    @patch("app.api.documents.get_document_metadata")
    @patch("app.api.documents.download_document")
    def test_pipeline_successful_pdf_retrieval(self, mock_download, mock_get_meta):
        doc_id = str(uuid.uuid4())
        expected_storage_path = f"{self.user_id}/{doc_id}_alex_resume.pdf"

        mock_get_meta.return_value = {
            "id": doc_id,
            "user_id": self.user_id,
            "file_name": "alex_resume.pdf",
            "file_type": PDF_MIME,
            "file_size": len(self.valid_pdf_bytes),
            "storage_path": expected_storage_path,
        }
        mock_download.return_value = self.valid_pdf_bytes

        response = self.client.get(f"/api/v1/documents/{doc_id}")

        self.assertEqual(response.status_code, 200)
        data = response.json()

        self.assertEqual(data["document_id"], doc_id)
        self.assertEqual(data["file_name"], "alex_resume.pdf")
        self.assertEqual(data["file_type"], PDF_MIME)
        self.assertEqual(data["storage_path"], expected_storage_path)
        self.assertIn("Alex Morgan", data["extracted_text"])

        # Verify exact stored path passed to download
        mock_download.assert_called_once_with(expected_storage_path)

    # 2. Successful DOCX Retrieval
    @patch("app.api.documents.get_document_metadata")
    @patch("app.api.documents.download_document")
    def test_pipeline_successful_docx_retrieval(self, mock_download, mock_get_meta):
        doc_id = str(uuid.uuid4())
        expected_storage_path = f"{self.user_id}/{doc_id}_jordan_resume.docx"

        mock_get_meta.return_value = {
            "id": doc_id,
            "user_id": self.user_id,
            "file_name": "jordan_resume.docx",
            "file_type": DOCX_MIME,
            "file_size": len(self.valid_docx_bytes),
            "storage_path": expected_storage_path,
        }
        mock_download.return_value = self.valid_docx_bytes

        response = self.client.get(f"/api/v1/documents/{doc_id}")

        self.assertEqual(response.status_code, 200)
        data = response.json()

        self.assertEqual(data["document_id"], doc_id)
        self.assertEqual(data["file_name"], "jordan_resume.docx")
        self.assertIn("Jordan Taylor", data["extracted_text"])
        self.assertIn("CPACC", data["extracted_text"])

        mock_download.assert_called_once_with(expected_storage_path)

    # 3. Missing Metadata (Returns 404, does NOT attempt download)
    @patch("app.api.documents.get_document_metadata")
    @patch("app.api.documents.download_document")
    def test_pipeline_missing_metadata_skips_download(self, mock_download, mock_get_meta):
        mock_get_meta.return_value = None
        doc_id = str(uuid.uuid4())

        response = self.client.get(f"/api/v1/documents/{doc_id}")

        self.assertEqual(response.status_code, 404)
        # Confirm download was NEVER attempted
        self.assertFalse(mock_download.called)

    # 4. Storage Download Failure (Returns safe 502, does NOT attempt extraction)
    @patch("app.api.documents.get_document_metadata")
    @patch("app.api.documents.download_document")
    @patch("app.api.documents.extract_text_from_pdf")
    def test_pipeline_download_failure_skips_extraction(self, mock_extract, mock_download, mock_get_meta):
        doc_id = str(uuid.uuid4())
        mock_get_meta.return_value = {
            "id": doc_id,
            "user_id": self.user_id,
            "file_name": "resume.pdf",
            "storage_path": f"{self.user_id}/{doc_id}_resume.pdf",
        }
        mock_download.side_effect = StorageServiceError("Remote storage service temporarily unavailable")

        response = self.client.get(f"/api/v1/documents/{doc_id}")

        self.assertEqual(response.status_code, 502)
        # Confirm extraction was NEVER called
        self.assertFalse(mock_extract.called)

    # 5. Extraction Failure (Returns safe 422, does not leak traceback)
    @patch("app.api.documents.get_document_metadata")
    @patch("app.api.documents.download_document")
    @patch("app.api.documents.extract_text_from_pdf")
    def test_pipeline_extraction_failure_returns_safe_error(self, mock_extract, mock_download, mock_get_meta):
        doc_id = str(uuid.uuid4())
        mock_get_meta.return_value = {
            "id": doc_id,
            "user_id": self.user_id,
            "file_name": "resume.pdf",
            "storage_path": f"{self.user_id}/{doc_id}_resume.pdf",
        }
        mock_download.return_value = self.valid_pdf_bytes
        mock_extract.side_effect = DocumentProcessingError("Unreadable PDF object stream")

        response = self.client.get(f"/api/v1/documents/{doc_id}")

        self.assertEqual(response.status_code, 422)
        self.assertIn("Unreadable PDF object stream", response.json()["error"])

    # =========================================================================
    # Task 4: Pipeline Consistency Check
    # =========================================================================
    @patch("app.api.documents.upload_document")
    @patch("app.api.documents.save_document_metadata")
    @patch("app.api.documents.get_document_metadata")
    @patch("app.api.documents.download_document")
    def test_upload_then_retrieval_consistency(
        self, mock_download, mock_get_meta, mock_save_meta, mock_upload
    ):
        """
        Verify that fields emitted during upload perfectly correspond to
        fields expected and returned during retrieval.
        """
        fixed_doc_id = str(uuid.uuid4())
        fixed_storage_path = f"{self.user_id}/{fixed_doc_id}_alex_resume.pdf"

        mock_upload.return_value = {"storage_path": fixed_storage_path, "status": "uploaded"}
        mock_save_meta.return_value = {
            "id": fixed_doc_id,
            "user_id": self.user_id,
            "file_name": "alex_resume.pdf",
            "storage_path": fixed_storage_path,
            "file_type": PDF_MIME,
            "file_size": len(self.valid_pdf_bytes),
        }

        # 1. Execute upload
        upload_resp = self.client.post(
            "/api/v1/documents/upload",
            data={"user_id": self.user_id},
            files={"file": ("alex_resume.pdf", self.valid_pdf_bytes, PDF_MIME)},
        )
        self.assertEqual(upload_resp.status_code, 200)
        upload_data = upload_resp.json()

        # 2. Simulate retrieval using data returned from upload
        mock_get_meta.return_value = {
            "id": upload_data["document_id"],
            "user_id": self.user_id,
            "file_name": upload_data["file_name"],
            "storage_path": upload_data["storage_path"],
            "file_type": upload_data["file_type"],
            "file_size": upload_data["file_size"],
        }
        mock_download.return_value = self.valid_pdf_bytes

        retrieval_resp = self.client.get(f"/api/v1/documents/{upload_data['document_id']}")
        self.assertEqual(retrieval_resp.status_code, 200)
        retrieval_data = retrieval_resp.json()

        # 3. Assert full contract consistency
        self.assertEqual(upload_data["document_id"], retrieval_data["document_id"])
        self.assertEqual(upload_data["file_name"], retrieval_data["file_name"])
        self.assertEqual(upload_data["file_type"], retrieval_data["file_type"])
        self.assertEqual(upload_data["file_size"], retrieval_data["file_size"])
        self.assertEqual(upload_data["storage_path"], retrieval_data["storage_path"])
        self.assertEqual(upload_data["extracted_text"], retrieval_data["extracted_text"])

    # =========================================================================
    # Task 5: Guarded Live Supabase Smoke Test (Opt-in only)
    # =========================================================================
    @unittest.skipUnless(
        os.environ.get("RUN_LIVE_SUPABASE_TESTS") == "1",
        "Live Supabase smoke test is opt-in (set RUN_LIVE_SUPABASE_TESTS=1).",
    )
    def test_live_supabase_smoke_test(self):
        """
        Guarded live smoke test that only executes if explicitly requested via RUN_LIVE_SUPABASE_TESTS=1.
        Uses a designated temporary path and guarantees automated cleanup.
        """
        from app.db.supabase import get_supabase_client, is_supabase_configured
        from app.services.storage_service import delete_document, upload_document

        if not is_supabase_configured():
            self.skipTest("Live Supabase credentials not configured in environment.")

        client = get_supabase_client()
        if client is None:
            self.skipTest("Live Supabase client initialization failed.")

        test_path = f"_automated_smoke_tests/{uuid.uuid4()}_smoke.pdf"
        try:
            # Test live upload
            upload_res = upload_document(
                file_bytes=self.valid_pdf_bytes,
                storage_path=test_path,
                content_type=PDF_MIME,
            )
            self.assertEqual(upload_res["status"], "uploaded")
        finally:
            # Guaranteed cleanup
            try:
                delete_document(test_path)
            except Exception:
                pass


if __name__ == "__main__":
    unittest.main()
